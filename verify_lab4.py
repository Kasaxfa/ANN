"""Сверка пересчитанного Excel с эталонными предсказаниями каждой строки."""
import json
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

out=Path(__file__).resolve().parent/'lab4_reworked'
metrics=json.loads((out/'metrics.json').read_text(encoding='utf-8'))
pred=pd.read_csv(out/'predictions.csv').set_index('source_csv_row')
book=load_workbook(out/'water_variant17_lab4_work.xlsx',read_only=True,data_only=True)
results={}
for activation in ('ReLU','Tanh'):
    for variant in ('baseline','tuned'):
        name=activation+'_'+variant
        mismatches=0
        error_rows=[]
        max_delta=0
        for row in book[name].iter_rows(min_row=2,values_only=True):
            source,target=row[:2]
            logit,prediction,error=row[30],row[32],row[33]
            assert isinstance(logit,(int,float)), 'Workbook needs Excel recalculation'
            expected=pred.loc[source]
            assert target==expected.target
            mismatches += int(prediction!=expected[name+'_pred'])
            assert error==int(prediction!=target)
            if error: error_rows.append(source)
            max_delta=max(max_delta,abs(logit-expected[name+'_logit']))
        expected_rows=pred[(pred.split=='test') & pred[name+'_error']].index.tolist()
        assert sorted(expected_rows)==sorted(error_rows)
        assert mismatches==0
        assert len(error_rows)==metrics['models'][activation][variant]['test']['errors']
        results[name]={'rows_checked':len(pred[pred.split=='test']),'prediction_mismatches':mismatches,
            'errors':len(error_rows),'error_source_csv_rows':error_rows,'max_logit_difference':max_delta}
book.close()
(out/'excel_verification.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
print(json.dumps(results,indent=2))
