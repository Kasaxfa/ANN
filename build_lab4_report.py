"""Курсовой отчёт по структуре предоставленного образца, из фактических метрик."""
from __future__ import annotations
import argparse
import json
import sys
from copy import deepcopy
from pathlib import Path
import pandas as pd
import plotly.graph_objects as go
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'lab4_reworked'
ASSETS=OUT/'report_assets'
REPORT=OUT/'Отчёт курсовая — ЛР4, вариант 17.docx'
M=json.loads((OUT/'metrics.json').read_text(encoding='utf-8'))
V=json.loads((OUT/'excel_verification.json').read_text(encoding='utf-8'))
def num(value,digits=4): return f'{value:.{digits}f}'.replace('.',',')
def pct(value): return num(value*100,4)+'%'

def plots():
    ASSETS.mkdir(exist_ok=True)
    hist=pd.read_csv(OUT/'training_metrics.csv')
    fig=go.Figure()
    for act in M['models']:
        part=hist[hist.activation==act]
        for split,dash in [('train','solid'),('validation','dash')]:
            fig.add_scatter(x=part.epoch,y=part[split+'_loss'],name=act+' / '+split,line={'dash':dash})
    fig.update_layout(title='Обучение базовых моделей без категорий',xaxis_title='Эпоха',yaxis_title='BCE loss')
    figures={'training':fig}
    fig=go.Figure()
    for stage,label in [('baseline','До подбора'),('tuned','После подбора')]:
        fig.add_bar(x=['ReLU','Tanh'],y=[M['models'][a][stage]['test']['errors'] for a in ['ReLU','Tanh']],name=label,
            text=[M['models'][a][stage]['test']['errors'] for a in ['ReLU','Tanh']],textposition='outside')
    fig.update_layout(title='Ошибки на независимой тестовой выборке (5 933 объекта)',yaxis_title='Число ошибок',barmode='group')
    figures['test_errors']=fig
    labels=['Пол','Активность','Погода','Пол + активность','Пол + погода','Активность + погода','Все три']
    fig=go.Figure()
    for act in M['models']:
        rows=[r for r in M['interactions'] if r['activation']==act]
        fig.add_bar(x=labels,y=[r['validation']['errors'] for r in rows],name=act)
    fig.update_layout(title='Раздельный и совместный подбор: ошибки validation',barmode='group',yaxis_title='Число ошибок',xaxis_tickangle=-18)
    figures['interactions']=fig
    fig=go.Figure()
    for act in M['models']:
        fig.add_bar(x=[value for field,value in M['category_keys']],y=M['models'][act]['codes'],name=act)
    fig.update_layout(title='Выбранные цифровые коды категорий',yaxis_title='Код — добавка к предактивациям',barmode='group')
    figures['codes']=fig
    for name,fig in figures.items():
        fig.update_layout(template='plotly_white',font={'family':'Arial','size':17},legend={'orientation':'h','y':1.12},margin={'l':75,'r':30,'b':100,'t':125})
        fig.write_html(ASSETS/(name+'.html'),include_plotlyjs=True)
        fig.write_image(ASSETS/(name+'.png'),width=1300,height=700,scale=1)

def build():
    source=ROOT.parent/'Отчёт курсовая.docx'
    doc=Document(source)
    # Сохраняем титульный лист, две авторские таблицы и стили образца.
    cut=doc.paragraphs[19]._p
    remove=False
    for node in list(doc._element.body):
        if node is cut: remove=True
        if remove and node.tag!=qn('w:sectPr'): doc._element.body.remove(node)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if 'Д.Ю. Михайлов' in p.text:
                        updated=p.text.replace('Д.Ю. Михайлов','В.А. Шамурзанов')
                        if p.runs:
                            p.runs[0].text=updated
                            for run in p.runs[1:]: run.text=''
                        else: p.add_run(updated)
    sec=doc.sections[-1]
    sec.page_width,sec.page_height=Cm(21),Cm(29.7)
    sec.top_margin,sec.bottom_margin,sec.left_margin,sec.right_margin=Cm(2),Cm(2),Cm(2),Cm(1)
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman'; normal.font.size=Pt(14)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.first_line_indent=Cm(1.25)
    for style in ('Heading 1','Heading 2'):
        st=doc.styles[style]
        st.font.name='Times New Roman'; st.font.size=Pt(14); st.font.bold=True
        st.paragraph_format.first_line_indent=Cm(0)
        st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.keep_with_next=True
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.clear()
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
    sec.different_first_page_header_footer=True

    def p(text):
        para=doc.add_paragraph(text)
        para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return para
    def h(text,level=1): return doc.add_heading(text,level)
    def listing(items):
        for i,text in enumerate(items,1): p(f'{i}. {text}')
    def table(caption,headers,rows):
        cap=p(caption);cap.paragraph_format.first_line_indent=Cm(0);cap.paragraph_format.keep_with_next=True
        tab=doc.add_table(rows=1,cols=len(headers));tab.style='Table Grid'
        for c,t in zip(tab.rows[0].cells,headers): c.text=str(t)
        repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for c,t in zip(tab.add_row().cells,values):c.text=str(t)
        for ri,row in enumerate(tab.rows):
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.first_line_indent=Cm(0)
                    para.paragraph_format.line_spacing=1
                    para.paragraph_format.space_after=Pt(4)
                    for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(11);run.bold=(ri==0)
        return tab
    def figure(name,caption):
        para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0)
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(ASSETS/(name+'.png')),width=Cm(17))
        para=p(caption);para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Cm(0)
    def code(text):
        para=doc.add_paragraph()
        para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1
        for i,line in enumerate(text.splitlines()):
            run=para.add_run(('\n' if i else '')+line);run.font.name='Consolas';run.font.size=Pt(9)

    h('ВВЕДЕНИЕ')
    p('В задачах машинного обучения используются числовые и категориальные признаки. Способ цифрового представления категорий влияет на результат классификации, поэтому выбор кодировки представляет самостоятельную практическую задачу [1, 2].')
    p('В курсовом проекте исследуется возврат категорий в уже обученную сеть. Сначала модель обучается на числовых признаках, затем её рабочий ход воспроизводится в Excel. Вводится дополнительный вход с единичными весами и подбираются значения категорий; исходные веса сети сохраняются.')
    p('Использован файл Daily_Water_Intake.csv, вариант 17. Целевая переменная — Hydration Level (Poor, Good); возвращаемые категории — пол, активность и погода. Методика основана на задании «Кодировка категорий» и предоставленном Excel-примере [5–7].')
    p('Цель курсового проекта — исследовать возможность уменьшения числа ошибок бинарной классификации за счёт цифровой кодировки категориальных признаков при фиксированных весах нейронной сети.')
    p('Для достижения цели в работе необходимо:')
    listing(['Подготовить данные и обучить ReLU и Tanh без категорий.',
        'Перенести рабочий ход в Excel и сверить предсказания.',
        'Ввести дополнительный вход и подобрать коды категорий.',
        'Исследовать отдельное и совместное влияние категорий.',
        'Оценить качество на независимой тестовой выборке.'])
    p('Подбор коэффициентов автоматизирован на Python. Значения перенесены в редактируемые ячейки Excel для ручного исследования. Отчёт описывает фактически выполненный программный поиск.')

    h('1. ПОСТАНОВКА ЗАДАЧИ')
    p('Решается задача бинарной классификации уровня гидратации по данным варианта 17. Нулевая метка соответствует Poor, единичная — Good. Для базовой сети используются три числовых признака: Age, Weight (kg), Daily Water Intake (liters). Три категориальных признака полностью исключаются из обучения и сохраняются отдельно для последующего эксперимента.')
    p('Обученная модель имеет структуру 3 → 10 → 1. Для включения категорий без повторного обучения формируется величина q — сумма кодов пола, активности и погоды. Эта величина поступает на каждый из десяти скрытых нейронов с фиксированным весом 1. Число скрытых и выходных нейронов не изменяется; число входов расширяется с трёх до четырёх.')
    p('Основной критерий подбора — число ошибок на проверочной выборке. Accuracy, Precision, Recall и F1 используются для сопоставления моделей. Порог вероятности равен 0,5. Условие logit ≥ 0 эквивалентно отнесению объекта к классу Good.')
    p('Качество оценивается раздельно на обучающей, проверочной и тестовой частях. Проверочная выборка используется для выбора эпохи и кодов; тестовая — для итоговой оценки. Такой порядок отделяет подбор параметров от проверки результата [4].')

    h('2. ОПИСАНИЕ ИСХОДНЫХ ДАННЫХ')
    h('2.1. Структура датасета',2)
    p(f"Исходный CSV-файл содержит {M['raw_rows']} строк и 7 столбцов. Строк с пропущенными значениями обнаружено {M['missing_rows']}. Удалено {M['duplicates_removed']} полных дубликатов; после очистки осталось {M['rows']} наблюдения. Номера исходных строк сохранены в поле source_csv_row, что позволяет проверить местоположение каждой ошибки.")
    p('Целевая переменная закодирована следующим образом: Poor = 0, Good = 1. После очистки класс Poor содержит 5 932 объекта, класс Good — 23 730 объектов. Большинство наблюдений относится к классу Good; простое предсказание этого класса даёт около 80% правильных ответов.')
    h('2.2. Числовые и категориальные признаки',2)
    table('Таблица 1 — Признаки набора данных',['Признак','Содержание','Использование'],[
        ['Age','Возраст','Обучение базовой сети'],['Weight (kg)','Масса тела','Обучение базовой сети'],
        ['Daily Water Intake (liters)','Потребление воды','Обучение базовой сети'],
        ['Gender','Female, Male','Подбор 2 кодов'],['Physical Activity Level','Low, Moderate, High','Подбор 3 кодов'],
        ['Weather','Cold, Normal, Hot','Подбор 3 кодов'],['Hydration Level','Poor, Good','Целевая переменная']])
    p('Масса тела сохранена как числовой признак. Ограничение на четыре выбранных признака из предыдущей лабораторной работы здесь не применяется: задание №4 требует исключить категории, а не уменьшить число числовых входов.')
    h('2.3. Разделение и нормализация',2)
    table('Таблица 2 — Состав выборок',['Выборка','Объектов','Poor','Good'],[[n,M['splits'][k]['n'],*M['splits'][k]['class_counts']] for k,n in [('train','Обучающая (60%)'),('validation','Проверочная (20%)'),('test','Тестовая (20%)')]])
    p('Разбиение стратифицировано, random_state = 42. Вначале выделяется 20% тестовых данных, затем четверть оставшейся части отводится под validation. Средние значения и стандартные отклонения рассчитываются только на обучающей части с помощью StandardScaler [4].')
    table('Таблица 3 — Параметры стандартизации',['Признак','Среднее','Стандартное отклонение'],[[n,num(a),num(b)] for n,a,b in zip(M['numeric_features'],M['normalization_mean'],M['normalization_std'])])
    p('Вычисление нормализованного входа: x = (v − μ) / σ. Перед обучением входы приведены к float32; эти же значения экспортированы в Excel. Категориальные коды добавляются к предактивациям после нормализации числовых данных и отдельно не стандартизируются.')

    h('3. БАЗОВАЯ НЕЙРОННАЯ СЕТЬ')
    h('3.1. Архитектура сети',2)
    p('Построены две модели с одинаковой структурой: три входа, скрытый слой из десяти нейронов и один выходной нейрон. Модели различаются функцией активации скрытого слоя — ReLU либо Tanh. Выходной слой формирует logit; sigmoid применяется при вычислении вероятности [3].')
    p('Прямой проход записывается в виде z = Wx + b; h = φ(z); l = vᵀh + b₀; p = 1 / (1 + exp(−l)). Для ReLU используется φ(z) = max(0, z), для Tanh — гиперболический тангенс.')
    h('3.2. Гиперпараметры обучения',2)
    table('Таблица 4 — Параметры обучения',['Параметр','Значение'],[
        ['Функция потерь','BCEWithLogitsLoss'],['Оптимизатор','Adam'],['Learning rate','0,001'],['Batch size','256'],
        ['Максимальное число эпох','150'],['Patience','20 эпох'],['Критерий сохранения','Минимальный validation loss'],['Seed','42']])
    p('Обе модели обучались независимо. Категориальные столбцы не передавались в обучающий набор. Веса сохранялись при уменьшении проверочной функции потерь более чем на 10⁻⁶; тестовые показатели не использовались для выбора эпохи.')
    p('Для обеих моделей лучшим оказалось состояние на 150-й эпохе. Обучение завершилось по ограничению числа эпох, а не по срабатыванию ранней остановки. Это означает, что полученные модели являются результатом заданного вычислительного эксперимента, но не доказанным оптимумом обучения.')
    h('3.3. Код модели (Python/PyTorch)',2)
    code('class WaterNet(nn.Module):\n    def __init__(self, activation):\n        super().__init__()\n        self.hidden = nn.Linear(3, 10)\n        self.activation = nn.ReLU() if activation == "ReLU" else nn.Tanh()\n        self.output = nn.Linear(10, 1)\n    def forward(self, x):\n        return self.output(self.activation(self.hidden(x))).squeeze(1)')
    h('3.4. Результаты обучения',2)
    table('Таблица 5 — Базовые модели на тестовой выборке',['Метрика','ReLU','Tanh'],[[key,pct(M['models']['ReLU']['baseline']['test'][key]),pct(M['models']['Tanh']['baseline']['test'][key])] for key in ['accuracy','precision','recall','f1']]+[['Ошибок',3,7]])
    p('ReLU допустила 3 ошибки, Tanh — 7 ошибок. Все ошибки исходных моделей на тесте относятся к ложноположительным: объекты Poor классифицированы как Good. Базовые сети уже обеспечивают высокую точность, поэтому ожидаемый резерв улучшения невелик.')
    figure('training','Рисунок 1 — Изменение функции потерь при обучении ReLU и Tanh')

    h('4. РЕАЛИЗАЦИЯ В EXCEL')
    h('4.1. Структура файлов',2)
    p('Создана одна рабочая книга water_variant17_lab4_work.xlsx. В ней объединены исходные данные, веса обеих сетей, редактируемые коды и четыре варианта прямого прохода. Расчётные листы содержат всю тестовую выборку — 5 933 объекта, без усечения до первых строк.')
    table('Таблица 6 — Основные листы рабочей книги',['Лист','Содержание'],[
        ['Summary','Сводные формулы ошибок, accuracy и проверки'],['Raw_data','Все 29 662 очищенные строки и принадлежность к выборке'],
        ['Input','Тестовые строки, категории и нормализованные числовые входы'],['Normalization','Средние и масштабы обучающей части'],
        ['Codes','Восемь значений категорий для каждой модели'],['Weights_ReLU / Weights_Tanh','Весовые матрицы, смещения и единичная строка'],
        ['ReLU_baseline / Tanh_baseline','Прямой проход с нулевым дополнительным входом'],['ReLU_tuned / Tanh_tuned','Прямой проход с выбранными кодами'],
        ['Interactions','Результаты раздельного и совместного подбора']])
    h('4.2. Воспроизведение рабочего хода',2)
    p('Числовой вход умножается на веса первого слоя с добавлением смещений. Затем применяются MAX(0, z) либо TANH(z), SUMPRODUCT для выходного слоя, sigmoid и пороговое правило. Параметры первого слоя экспортированы без округления, чтобы сохранить результат вычисления.')
    p('В транспонированную матрицу первого слоя добавлена строка из десяти единиц. Во входной вектор добавлен столбец q, изначально равный нулю. При таком расширении вклад новых связей равен нулю, поэтому исходные logit, метки и ошибки сохраняются. После подстановки кодов формула принимает вид zⱼ = bⱼ + Σᵢwⱼᵢxᵢ + q.')
    code('q = code(Gender) + code(Activity) + code(Weather)\nz_j = bias_j + x_age*w_age_j + x_weight*w_weight_j + x_water*w_water_j + q\nh_j = MAX(0, z_j)  или  TANH(z_j)\nlogit = SUMPRODUCT(h_1:h_10, output_weights) + output_bias\nprediction = IF(logit >= 0, 1, 0)\nerror = IF(prediction <> target, 1, 0)')
    h('4.3. Проверка правильности расчёта',2)
    p('Рабочая книга открыта и полностью пересчитана Microsoft Excel. После сохранения числовые значения формул прочитаны обратно и сопоставлены с предсказаниями Python для каждой исходной строки CSV. Проверялись не только суммы ошибок, но и их номера.')
    table('Таблица 7 — Сверка Excel и Python',['Вариант','Проверено строк','Расхождений меток','Ошибок'],[[n,v['rows_checked'],v['prediction_mismatches'],v['errors']] for n,v in V.items()])
    p('Количество и местоположение ошибок совпали для всех четырёх вариантов. Максимальное абсолютное различие logit между сохранёнными CSV-значениями Python и Excel составляет менее 10⁻¹². Хеши весов до и после подбора совпадают. Результаты проверки сохранены в excel_verification.json.')
    p('Для ручного исследования необходимо изменить жёлтые ячейки листа Codes и пересчитать книгу. Формулы tuned обновятся, а столбцы Python_prediction останутся эталоном первоначально найденного решения. Поэтому после ручных изменений отличия от этого эталона ожидаемы.')

    h('5. ПОДБОР КОЭФФИЦИЕНТОВ ДЛЯ ТРЁХ ПРИЗНАКОВ')
    h('5.1. Методика',2)
    p('Для пола задаются два коэффициента, для активности и погоды — по три, всего восемь значений. По строке данных выбирается одно значение каждого признака; их сумма образует дополнительный вход q. Этот способ реализует единичную строку весов из методического примера [6, 7].')
    p('Подбор выполнен дифференциальной эволюцией [8] в диапазоне от −2 до 2. Использованы seed = 42, до 45 поколений и множитель размера популяции 6. Затем проведено покоординатное уточнение с шагами 0,1; 0,02; 0,005. Экспортируемые коэффициенты округлены до трёх десятичных знаков; качество пересчитано именно с этими значениями.')
    p('Целевая функция равна числу ошибок validation с малой добавкой 10⁻⁶Σc². Эта добавка лишь разрешает равенство числа ошибок в пользу меньших коэффициентов. Нулевая кодировка включена в сравнение, поэтому ухудшающая validation настройка не принимается.')
    p('Исследованы семь наборов допустимых изменений: каждый признак отдельно, каждая пара и все три одновременно. Для каждой модели выбиралась первая настройка с минимальным числом проверочных ошибок в порядке этого перебора. Это локальный вычислительный поиск; глобальная оптимальность найденных коэффициентов не утверждается.')
    h('5.2. Итоговые коэффициенты',2)
    table('Таблица 8 — Выбранные цифровые коды',['Признак','Категория','ReLU','Tanh'],[[f,v,num(M['models']['ReLU']['codes'][i],3),num(M['models']['Tanh']['codes'][i],3)] for i,(f,v) in enumerate(M['category_keys'])])
    p('Для Tanh выбрана настройка только признака Gender: Female = −0,011, Male = −0,015. Коды активности и погоды равны нулю. Она уменьшила число ошибок validation с 11 до 0 и число тестовых ошибок с 7 до 0.')
    p('Для ReLU выбрана совместная настройка пола и погоды. Оба значения пола получили код −0,007, то есть эта составляющая является общей добавкой и не различает категории. Погода задаёт коды Cold = −0,012, Normal = −0,015, Hot = 0. Активность не используется. Ошибки validation уменьшились с 6 до 0, но число тестовых ошибок выросло с 3 до 4.')
    p('Одинаковые или близкие коды нельзя трактовать как доказательство самостоятельной информативности категорий. В частности, эффект Gender у Tanh может преимущественно отражать небольшую общую корректировку предактиваций. В данной схеме существенна сумма трёх кодов; отдельные коэффициенты не имеют единственной интерпретации.')
    figure('codes','Рисунок 2 — Значения выбранных кодов для ReLU и Tanh')
    h('5.3. Взаимовлияние категориальных столбцов',2)
    left=[r for r in M['interactions'] if r['activation']=='ReLU']
    right=[r for r in M['interactions'] if r['activation']=='Tanh']
    field_labels={'Gender':'Пол','Physical Activity Level':'Активность','Weather':'Погода'}
    table('Таблица 9 — Раздельный и совместный подбор',['Изменяемые признаки','ReLU val / test','Tanh val / test'],[[' + '.join(field_labels[f] for f in a['fields']),f"{a['validation']['errors']} / {a['test']['errors']}",f"{b['validation']['errors']} / {b['test']['errors']}"] for a,b in zip(left,right)])
    p('В таблице приведены ошибки на проверочной и тестовой выборках. Каждая строка соответствует самостоятельному подбору при фиксированных остальных кодах, равных нулю. Влияние пары не обязано совпадать с суммой отдельных эффектов: складывающиеся коды проходят через нелинейную функцию активации и порог классификации.')
    p('Результаты строк таблицы используются для описания эксперимента. Итоговая кодировка выбиралась по validation; дополнительные тестовые результаты не использовались для повторного подбора коэффициентов.')
    figure('interactions','Рисунок 3 — Число ошибок validation при разных сочетаниях категориальных признаков')

    h('6. РЕЗУЛЬТАТЫ')
    h('6.1. Метрики качества классификации',2)
    p('Accuracy — доля правильных ответов. Precision — доля истинных объектов Good среди предсказанных Good. Recall — доля найденных объектов Good среди всех объектов этого класса. F1 — гармоническое среднее Precision и Recall. Строки матрицы ошибок соответствуют истинным классам, столбцы — предсказанным, порядок классов: Poor, Good.')
    h('6.2. Результаты на тестовой выборке',2)
    variants=[(a,s) for a in ['ReLU','Tanh'] for s in ['baseline','tuned']]
    headers=['Метрика','ReLU до','ReLU после','Tanh до','Tanh после']
    rows=[[k,*[pct(M['models'][a][s]['test'][k]) for a,s in variants]] for k in ['accuracy','precision','recall','f1']]
    rows += [['Ошибок',*[M['models'][a][s]['test']['errors'] for a,s in variants]]]
    table('Таблица 10 — Метрики на 5 933 тестовых объектах',headers,rows)
    for a,s in variants:
        score=M['models'][a][s]['test'];cm=score['confusion_matrix']
        p(f"{a}, {'до' if s=='baseline' else 'после'} подбора: TN = {cm[0][0]}, FP = {cm[0][1]}, FN = {cm[1][0]}, TP = {cm[1][1]}. Accuracy = {pct(score['accuracy'])}; ошибок — {score['errors']}.")
    figure('test_errors','Рисунок 4 — Сравнение числа ошибок до и после подбора кодов')
    h('6.3. Результаты на полном наборе данных',2)
    table('Таблица 11 — Описательная оценка на всех очищенных данных',['Модель','До: ошибок / accuracy','После: ошибок / accuracy'],[[a,f"{m['all_baseline']['errors']} / {pct(m['all_baseline']['accuracy'])}",f"{m['all_tuned']['errors']} / {pct(m['all_tuned']['accuracy'])}"] for a,m in M['models'].items()])
    p('Полная выборка содержит обучающие и проверочные объекты, поэтому её показатели приведены как описательная характеристика. Независимой итоговой оценкой остаётся результат на выделенном тесте.')
    h('6.4. Сравнительный анализ и ограничения',2)
    delta=(M['models']['Tanh']['tuned']['test']['accuracy']-M['models']['Tanh']['baseline']['test']['accuracy'])*100
    p(f'У Tanh подбор устранил семь тестовых ошибок и повысил accuracy на {num(delta)} процентного пункта. У ReLU результат на тесте ухудшился на одну ошибку, несмотря на отсутствие ошибок validation. Следовательно, уменьшение проверочной ошибки само по себе не гарантирует улучшения на новых объектах.')
    p('Обе настроенные модели получили ноль ошибок validation. Правило выбора по этому показателю с сохранением первой модели при равенстве выбрало бы ReLU. Tanh имеет лучший наблюдаемый результат на тесте, но это не следует выдавать за результат выбора архитектуры на тестовых данных. В работе приведены обе заранее заданные архитектуры.')
    p('Базовые сети уже разделяют данные почти безошибочно. Поэтому влияние найденных кодов невелико и относится к нескольким пограничным объектам. Нулевое число ошибок Tanh в данном разбиении не означает гарантированную точность 100% на любых данных. Для вывода об устойчивости эффекта нужны повторные разбиения и независимые данные; в этом эксперименте они не исследовались.')
    p('Высокая точность характеризует воспроизведение меток предоставленного CSV. Работа не проверяет способ получения исходных меток и не устанавливает применимость модели для медицинской диагностики. Поиск кодов не доказывает причинного влияния пола, активности или погоды на гидратацию.')

    h('7. ЗАКЛЮЧЕНИЕ')
    p('В ходе курсового проекта разработана и исследована схема включения категориальных признаков в обученную нейронную сеть при фиксированных весах на данных Daily Water Intake, вариант 17.')
    p('Выполнены следующие задачи:')
    listing(['Подготовлены данные: удалены 338 дубликатов, сохранены 29 662 наблюдения, сформированы обучающая, проверочная и тестовая выборки.',
        'Обучены сети 3 → 10 → 1 с активациями ReLU и Tanh на трёх числовых признаках; категориальные признаки из обучения исключены.',
        'Реализован прямой проход в Excel. Для всех 5 933 тестовых объектов совпали предсказания и номера ошибок Python и электронной таблицы.',
        'Введён дополнительный вход с единичными весами. Подобраны коды пола, активности и погоды при неизменных исходных параметрах сетей.',
        'Проверено раздельное и совместное влияние трёх категориальных столбцов; построены графики обучения, коэффициентов и ошибок.',
        'Для Tanh число тестовых ошибок уменьшено с 7 до 0; Accuracy, Precision, Recall и F1 на данном тесте составили 100%. Для ReLU ошибки изменились с 3 до 4.',
        'Сохранены модели, рабочая книга Excel, таблицы предсказаний и метрик, результаты проверки и иллюстрации.'])
    p('В исследованном разбиении возможность улучшения классификации без переобучения исходных весов показана для Tanh. Результат ReLU демонстрирует ограничение метода: настройка по validation может ухудшить тестовое качество. Поэтому эффект кодировки необходимо проверять на данных, не использованных при её подборе.')
    p('Для дальнейшей проверки эффекта необходимы повторные разбиения и сравнение с постоянной добавкой к предактивациям.')

    h('СПИСОК ЛИТЕРАТУРЫ')
    refs=[
        'Goodfellow I., Bengio Y., Courville A. Deep Learning. MIT Press, 2016. https://www.deeplearningbook.org/',
        'Chollet F. Deep Learning with Python. Manning Publications, 2018.',
        'Документация PyTorch: https://pytorch.org/docs/',
        'Документация scikit-learn: https://scikit-learn.org/ ; Common pitfalls: https://scikit-learn.org/stable/common_pitfalls.html',
        'Daily Water Intake & Hydration Patterns Dataset. Предоставленный файл Daily_Water_Intake.csv, вариант 17. Исходные данные расчёта.',
        'Кодировка категорий. Методические указания к лабораторной работе. Файл «кодировкаКатегорий_26_02_15.docx», предоставленный с заданием.',
        'Пример реализации рабочего хода нейронной сети в электронных таблицах. Файл «lab_4_work_26_02_17.xlsx», предоставленный с заданием.',
        'Документация SciPy. scipy.optimize.differential_evolution: https://docs.scipy.org/doc/scipy/reference/generated/scipy.optimize.differential_evolution.html']
    for i,ref in enumerate(refs,1):p(f'{i}. {ref}')
    h('ПРИЛОЖЕНИЯ')
    h('Приложение А. Основной код расчёта',2)
    p('Файл lab4_recalculation.py. Модуль формирования Excel — lab4_workbook.py; сборка отчёта — build_lab4_report.py. Полные версии сохранены в каталоге проекта.')
    for line in (ROOT/'lab4_recalculation.py').read_text(encoding='utf-8').splitlines():
        para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1
        para.paragraph_format.space_after=Pt(0)
        run=para.add_run(line);run.font.name='Consolas';run.font.size=Pt(8)
    h('Приложение Б. Местоположение тестовых ошибок',2)
    for name,v in V.items(): p(name+': '+(', '.join(map(str,v['error_source_csv_rows'])) if v['error_source_csv_rows'] else 'ошибок нет')+'.')
    p('Номера соответствуют строкам исходного CSV с учётом строки заголовка. Полный список объектов, меток и предсказаний находится в predictions.csv.')
    h('Приложение В. Порядок воспроизведения',2)
    listing(['Запустить python lab4_recalculation.py --output-dir lab4_reworked.',
        'Открыть lab4_reworked/water_variant17_lab4_work.xlsx в Excel, выполнить полный пересчёт и сохранить книгу.',
        'Запустить python verify_lab4.py для построчной проверки.',
        'Запустить python build_lab4_report.py для графиков и отчёта.'])
    doc.core_properties.title='Курсовой проект: кодировка категорий, вариант 17'
    doc.core_properties.author='В.А. Шамурзанов'
    doc.save(REPORT)
    print(str(REPORT),flush=True)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    parser=argparse.ArgumentParser();parser.add_argument('--skip-plots',action='store_true')
    args=parser.parse_args()
    if not args.skip_plots: plots()
    build()
