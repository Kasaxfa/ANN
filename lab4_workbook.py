"""Excel-прямой проход с редактируемыми кодами категорий и контролем ошибок."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

def save_workbook(path, frame, x, y, splits, metrics, predictions):
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append(['Показатель','ReLU baseline','ReLU tuned','Tanh baseline','Tanh tuned'])
    summary.append(['Ошибки TEST'])
    summary.append(['Accuracy TEST'])
    summary.append(['Расхождения с Python (pred)'])
    summary.append(['Макс. отклонение logit'])
    summary.append(['Число объектов TEST', len(splits['test'])])
    summary.append(['Train / validation / test','60% / 20% / 20%'])
    summary.append(['Подбор кодов','Только validation; веса фиксированы'])
    summary.append(['Редактирование','Коды категорий: жёлтые ячейки листа Codes. Пересчитать Excel (F9).'])
    summary.append(['Проверка','При изменении кодов отличия tuned от исходных Python-предсказаний ожидаемы.'])
    summary.append(['Формула','z = b + W*x + q*1; q = code(Gender) + code(Activity) + code(Weather)'])
    summary.append(['Нулевая кодировка','При q=0 расширенная сеть совпадает с исходной.'])
    summary.append(['Источник строк','source_csv_row — строка исходного CSV с учётом заголовка.'])
    codes = book.create_sheet('Codes')
    codes.append(['Field','Category','ReLU','Tanh'])
    for i,(field,category) in enumerate(metrics['category_keys']):
        codes.append([field,category,metrics['models']['ReLU']['codes'][i],metrics['models']['Tanh']['codes'][i]])
    for row in codes.iter_rows(min_row=2,min_col=3,max_col=4):
        for cell in row: cell.fill=PatternFill('solid',fgColor='FFF2CC')
    raw = book.create_sheet('Raw_data')
    raw.append(list(frame.columns)+['split','target'])
    for i,row in enumerate(frame.itertuples(index=False,name=None)):
        raw.append(list(row)+[predictions.loc[i,'split'],int(y[i])])
    inp = book.create_sheet('Input')
    inp.append(['source_csv_row','split','target','Age','Weight (kg)','Water (liters)',
        'Gender','Activity','Weather','x_age','x_weight','x_water'])
    for idx in splits['test']:
        row=frame.iloc[idx]
        inp.append([int(row.source_csv_row),'test',int(y[idx]),float(row.Age),float(row['Weight (kg)']),
            float(row['Daily Water Intake (liters)']),row.Gender,row['Physical Activity Level'],row.Weather,*x[idx].tolist()])
    norm=book.create_sheet('Normalization')
    norm.append(['Feature','Train mean','Train std'])
    for values in zip(metrics['numeric_features'],metrics['normalization_mean'],metrics['normalization_std']): norm.append(list(values))
    table=book.create_sheet('Interactions')
    table.append(['Model','Allowed categories','Validation errors','Test errors','Codes'])
    for item in metrics['interactions']:
        table.append([item['activation'],' + '.join(item['fields']),item['validation']['errors'],item['test']['errors'],str(item['codes'])])
    last=len(splits['test'])+1
    for activation in ('ReLU','Tanh'):
        weights=metrics['models'][activation]['weights']
        ws=book.create_sheet('Weights_'+activation)
        ws.append(['Parameter']+['h'+str(j+1) for j in range(10)])
        ws.append(['bias']+weights['hidden_bias'])
        for j,feature in enumerate(metrics['numeric_features']):
            ws.append([feature]+[row[j] for row in weights['hidden_weight']])
        ws.append(['q (unit row)']+[1]*10)
        ws.append([])
        ws.append(['Output layer']+['h'+str(j+1) for j in range(10)])
        ws.append(['Weights']+weights['output_weight'])
        ws.append(['Bias',weights['output_bias']])
        for variant in ('baseline','tuned'):
            name=activation+'_'+variant
            calc=book.create_sheet(name)
            calc.append(['source_csv_row','target','zero','x_age','x_weight','x_water',
                'q_gender','q_activity','q_weather','q_sum']+
                [f'z{j}' for j in range(1,11)]+[f'h{j}' for j in range(1,11)]+
                ['logit','probability','prediction','error','Python_prediction','match','abs_logit_diff'])
            for pos,idx in enumerate(splits['test'],start=2):
                r=pos
                cells=[f'=Input!A{r}',f'=Input!C{r}',0,f'=Input!J{r}',f'=Input!K{r}',f'=Input!L{r}']
                for field,col in [('Gender','G'),('Physical Activity Level','H'),('Weather','I')]:
                    code_col='C' if activation=='ReLU' else 'D'
                    cells.append(0 if variant=='baseline' else f'=SUMIFS(Codes!${code_col}$2:${code_col}$9,Codes!$A$2:$A$9,"{field}",Codes!$B$2:$B$9,Input!{col}{r})')
                cells.append(f'=C{r}+SUM(G{r}:I{r})')
                for j in range(10):
                    c=get_column_letter(j+2)
                    cells.append(f'={ws.title}!{c}$2+D{r}*{ws.title}!{c}$3+E{r}*{ws.title}!{c}$4+F{r}*{ws.title}!{c}$5+J{r}*{ws.title}!{c}$6')
                for j in range(10):
                    c=get_column_letter(j+11)
                    cells.append(f'=MAX(0,{c}{r})' if activation=='ReLU' else f'=TANH({c}{r})')
                cells += [f'=SUMPRODUCT(U{r}:AD{r},{ws.title}!$B$9:$K$9)+{ws.title}!$B$10',
                    f'=1/(1+EXP(-AE{r}))',f'=IF(AE{r}>=0,1,0)',f'=IF(AG{r}<>B{r},1,0)',
                    int(predictions.loc[idx,name+'_pred']),f'=IF(AG{r}=AI{r},1,0)',
                    f'=ABS(AE{r}-({float(predictions.loc[idx,name+"_logit"]):.17g}))']
                calc.append(cells)
            sc=2+(0 if activation=='ReLU' else 2)+(0 if variant=='baseline' else 1)
            summary.cell(2,sc,f'=SUM({name}!AH2:AH{last})')
            summary.cell(3,sc,f'=1-{get_column_letter(sc)}2/$B$6')
            summary.cell(4,sc,f'=$B$6-SUM({name}!AJ2:AJ{last})')
            summary.cell(5,sc,f'=MAX({name}!AK2:AK{last})')
    for sheet in book:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='24476B')
            cell.alignment=Alignment(wrap_text=True)
        sheet.row_dimensions[1].height=32
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=18
    summary.column_dimensions['A'].width=36
    codes.column_dimensions['A'].width=30
    book.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True)
    book.save(path)
